"""
Escritura de los archivos que lee Power BI: parquet y Excel, siempre atómica.

PARA QUÉ SIRVE
--------------
Las tareas corren cada minuto y Power BI refresca cada 10 leyendo esos
mismos archivos. Dos reglas:

1. ATÓMICA: se escribe en un temporal y se reemplaza de una vez. Power BI
   nunca lee un archivo a medio escribir.
2. SI EL ARCHIVO ESTÁ TOMADO, SE SALTEA: alguien lo abrió en Excel, o Power
   BI lo está leyendo. Esa corrida termina con código 3 y un "[SKIP]" en el
   log; la del minuto siguiente escribe. Es un estado normal, no un error, y
   no debe dejar un traceback que asuste a quien lea el log.
"""
import os
import time

import openpyxl
from openpyxl.styles import Font

CODIGO_SALTEADO = 3


class ArchivoTomado(Exception):
    """El destino lo tiene abierto otro proceso; se escribe en la próxima corrida."""


def archivo_tomado(ruta) -> bool:
    """True si otro proceso (Excel, Power BI) tiene el archivo abierto.

    Se chequea ANTES de llamar a la API: sin esto, cada corrida gastaría
    20 segundos de API para morir en el reemplazo del final.
    """
    try:
        with open(ruta, "r+b"):
            return False
    except PermissionError:
        return True
    except FileNotFoundError:
        return False


def _reemplazar(temporal, destino, reintentos=5, espera=3):
    """os.replace con reintento corto; si sigue tomado, ArchivoTomado."""
    for intento in range(1, reintentos + 1):
        try:
            os.replace(temporal, destino)
            return
        except PermissionError:
            if intento == reintentos:
                os.remove(temporal)
                raise ArchivoTomado(os.path.basename(str(destino)))
            time.sleep(espera)


def escribir_parquet(df, ruta, **kwargs):
    kwargs.setdefault("index", False)
    temporal = f"{ruta}.tmp.{os.getpid()}"
    df.to_parquet(temporal, **kwargs)
    _reemplazar(temporal, ruta)


def escribir_excel(df, ruta, hoja="Datos", formatos=None):
    """Una hoja con encabezado en negrita y formato por columna.

    `formatos` = {columna: formato de número de Excel}. Los tipos importan:
    Power Query castea cada columna (fecha, entero) y una celda con el tipo
    equivocado hace fallar el refresh entero.
    """
    formatos = formatos or {}
    libro = openpyxl.Workbook()
    ws = libro.active
    ws.title = hoja
    ws.append(list(df.columns))
    for celda in ws[1]:
        celda.font = Font(bold=True)
    for fila in df.itertuples(index=False):
        ws.append([None if v != v else v for v in fila])          # NaN -> celda vacía
    for j, col in enumerate(df.columns, start=1):
        if col in formatos:
            for (celda,) in ws.iter_rows(min_row=2, min_col=j, max_col=j):
                celda.number_format = formatos[col]
    ws.freeze_panes = "A2"
    temporal = f"{ruta}.tmp.{os.getpid()}.xlsx"
    libro.save(temporal)
    _reemplazar(temporal, ruta)
